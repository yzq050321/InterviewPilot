"""Resume text extraction with explicit size/page/type limits.

Only extracted text is returned to the caller. The original file is never
persisted by this module, which keeps the V2.2 resume flow easier to explain
and reduces the amount of candidate PII stored by the application.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
import zipfile

from docx import Document
from pypdf import PdfReader


MAX_RESUME_BYTES = 5 * 1024 * 1024
MAX_RESUME_PAGES = 20
MAX_RESUME_CHARS = 8000
MIN_RESUME_CHARS = 30
MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024

ALLOWED_SUFFIXES = {".pdf", ".docx"}


class ResumeParseError(ValueError):
    """A user-facing validation error; message is safe to return to the client."""


@dataclass(frozen=True)
class ParsedResume:
    filename: str
    text: str
    pages: int | None


def _normalize_text(value: str) -> str:
    # Normalize whitespace but keep line breaks so prompt context stays readable.
    value = value.replace("\x00", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _safe_filename(filename: str | None) -> str:
    # Strip directories and cap the stored display name.
    # 必须同时按 "/" 和 "\\" 切分：Linux 上的 Path(...).name 只把 "/" 当作
    # 分隔符，而浏览器/Windows 客户端可能传来反斜杠路径。只依赖 Path 会让
    # 反斜杠路径在 Linux 服务器上原样保留，等于没剥离目录。
    raw = (filename or "").replace("\\", "/")
    name = raw.rsplit("/", 1)[-1].strip()
    if not name or name in {".", ".."}:
        name = "resume"
    return name[:255]


def extract_resume_text(filename: str | None, data: bytes) -> ParsedResume:
    """Extract plain text from a PDF/DOCX resume.

    Raises ResumeParseError for unsupported, empty, oversized, encrypted or
    scanned documents. The caller maps that to HTTP 400/413.
    """
    safe_name = _safe_filename(filename)
    suffix = Path(safe_name).suffix.lower()

    if suffix not in ALLOWED_SUFFIXES:
        raise ResumeParseError("仅支持 PDF 或 DOCX 格式的简历")
    if not data:
        raise ResumeParseError("简历文件为空")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeParseError("简历文件不能超过 5 MB")

    pages: int | None = None

    if suffix == ".pdf":
        # PDF header is allowed to appear after a small amount of leading data.
        if b"%PDF-" not in data[:1024]:
            raise ResumeParseError("文件内容不是有效的 PDF")
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ResumeParseError("暂不支持加密 PDF，请导出为普通 PDF")
            pages = len(reader.pages)
            if pages > MAX_RESUME_PAGES:
                raise ResumeParseError(f"简历不能超过 {MAX_RESUME_PAGES} 页")
            parts = [page.extract_text() or "" for page in reader.pages]
        except ResumeParseError:
            raise
        except Exception as exc:
            raise ResumeParseError("PDF 解析失败，请确认文件未损坏") from exc
    else:
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ResumeParseError("文件内容不是有效的 DOCX")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "word/document.xml" not in names:
                    raise ResumeParseError("文件内容不是有效的 DOCX")
                uncompressed = sum(item.file_size for item in archive.infolist())
                if uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise ResumeParseError("DOCX 解压后体积异常，已拒绝处理")
            document = Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        except ResumeParseError:
            raise
        except Exception as exc:
            raise ResumeParseError("DOCX 解析失败，请确认文件未损坏") from exc

    text = _normalize_text("\n".join(parts))
    if len(text) < MIN_RESUME_CHARS:
        raise ResumeParseError("未提取到足够的简历文字；扫描版 PDF/图片简历暂不支持")
    if len(text) > MAX_RESUME_CHARS:
        text = text[:MAX_RESUME_CHARS].rstrip()

    return ParsedResume(filename=safe_name, text=text, pages=pages)
